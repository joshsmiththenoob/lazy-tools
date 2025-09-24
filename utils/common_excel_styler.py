# utils/common_excel_styler.py
from typing import Protocol, Dict, Callable, Union
from openpyxl.styles import PatternFill
from openpyxl.worksheet.worksheet import Worksheet

# 共用顏色
FILL_YELLOW = PatternFill(start_color="FFF5CC", end_color="FFF5CC", fill_type="solid")
FILL_RED    = PatternFill(start_color="FFECEC", end_color="FFECEC", fill_type="solid")
FILL_GREEN  = PatternFill(start_color="E8F5E9", end_color="E8F5E9", fill_type="solid")
FILL_ORANGE = PatternFill(start_color="FFA500", end_color="FFA500", fill_type="solid")
FILL_BLUE   = PatternFill(start_color="E3F2FD", end_color="E3F2FD", fill_type="solid")
FILL_GRAY   = PatternFill(start_color="EEEEEE", end_color="EEEEEE", fill_type="solid")

class WorksheetStyler(Protocol):
    """各專案的 Styler 只要實作 apply(ws) 即可"""
    def apply(self, ws: Worksheet) -> None: ...

def header_index_map(ws: Worksheet) -> Dict[str, int]:
    """回傳：{欄名 -> 欄號}，方便用欄名取到第幾欄"""
    return {ws.cell(row=1, column=c).value: c for c in range(1, ws.max_column + 1)}

def color_if(ws: Worksheet, col_name: str, predicate: Callable[[Union[float, int, str, None]], bool], fill: PatternFill):
    """
    依條件替特定欄上色
    
    針對特定Work sheet，在特定欄位上，以自訂的規則(ex: 匿名函數) 來決定是否上色
    """
    headers = header_index_map(ws)
    if col_name not in headers:
        return
    c_idx = headers[col_name]
    for r in range(2, ws.max_row + 1):
        cell = ws.cell(row=r, column=c_idx)
        val = cell.value
        # 嘗試把字串數值轉成 float
        if isinstance(val, str):
            try:
                val = float(val.replace(",", ""))
            except Exception:
                pass
        if predicate(val):
            cell.fill = fill
