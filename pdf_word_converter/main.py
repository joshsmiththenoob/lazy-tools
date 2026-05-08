# main.py
from __future__ import annotations

from pathlib import Path
from typing import List
from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn


PROJECT_DIR = Path(__file__).resolve().parent
INPUT_DIR = PROJECT_DIR / "input"
OUTPUT_DIR = PROJECT_DIR / "output"


def find_first_pdf(input_dir: Path) -> Path:
    pdfs = sorted(input_dir.glob("*.pdf"))
    if not pdfs:
        raise FileNotFoundError(f"No PDF found in: {input_dir}")
    return pdfs[0]


def set_default_chinese_font(doc: Document, font_name: str = "標楷體", font_size_pt: int = 12) -> None:
    """
    設定 Word 預設字型（包含東亞字型），避免中文變成奇怪的字體。
    """
    style = doc.styles["Normal"]
    style.font.name = font_name
    style.font.size = Pt(font_size_pt)
    # East Asian font
    style.element.rPr.rFonts.set(qn("w:eastAsia"), font_name)


def build_last_page_form_docx(
    out_path: Path,
    headers: List[str],
    blank_rows: int = 12,
) -> None:
    doc = Document()
    set_default_chinese_font(doc, font_name="標楷體", font_size_pt=12)

    # Title (match PDF last page meaning)
    p1 = doc.add_paragraph("附件")
    p1.alignment = WD_ALIGN_PARAGRAPH.LEFT

    p2 = doc.add_paragraph("正體中文文本語料資料集提報清單")
    p2.runs[0].bold = True
    p2.runs[0].font.size = Pt(16)
    p2.alignment = WD_ALIGN_PARAGRAPH.LEFT

    doc.add_paragraph("")  # spacer

    # Table
    n_cols = len(headers)
    table = doc.add_table(rows=1 + blank_rows, cols=n_cols)
    table.style = "Table Grid"

    # Header row
    for j, h in enumerate(headers):
        cell = table.cell(0, j)
        cell.text = h
        # Make header bold
        for run in cell.paragraphs[0].runs:
            run.bold = True

    # Blank rows
    for i in range(1, 1 + blank_rows):
        for j in range(n_cols):
            table.cell(i, j).text = ""  # keep editable empty cells

    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    doc.save(out_path)


def main() -> None:
    pdf_path = find_first_pdf(INPUT_DIR)

    # 依你這份 PDF 最後一頁的表格欄位（提供機關 / 資料集名稱 / 資料集連結 / 備註）
    headers = ["提供機關", "資料集名稱", "資料集連結", "備註"]

    out_docx = OUTPUT_DIR / "last_page_form.docx"
    build_last_page_form_docx(out_docx, headers=headers, blank_rows=12)

    print(f"[OK] Input PDF: {pdf_path.name}")
    print(f"[OK] Output DOCX: {out_docx}")


if __name__ == "__main__":
    main()
