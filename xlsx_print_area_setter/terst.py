# excel_postprocessor.py
from pathlib import Path
from openpyxl import load_workbook
from openpyxl.worksheet.pagebreak import Break
from openpyxl.worksheet.properties import PageSetupProperties

class DailyReportPostProcessor:
    """
    目的：
      - 清掉所有舊的列印/分頁設定
      - 列印範圍設成 A1:U{last_row}
      - 動態尋找 '花卉(把)'，把它放在第 2 頁第一列
      - 寬度 1 頁、高度不限（保留手動分頁）
    """

    # 你可以依資料格式調整這些預設
    LEFT_COL = "A"
    RIGHT_COL = "U"
    REPEAT_TITLE_ROWS = "1:8"       # 每頁重複列印的表頭（不要就設 None）
    TARGET_TEXT = "花卉(把)"        # 要成為第二頁第一列的關鍵字
    SEARCH_MAX_COL = 21             # A..U → 21
    SEARCH_START_ROW = 1            # 從第幾列開始掃描關鍵字
    PRINT_LAST_ROW_FALLBACK = 147   # 若找不到實際最後列，至少印到這一列
    DATA_ROW_START = 9
    DATA_ROW_END_FALLBACK = 152  # 覆蓋到說明區，之後會用實際最後列覆蓋
    OL_M, COL_S = 13, 19      # M~S
    
    def __init__(self, file_path: str):
        self.file_path = str(file_path)

    # -------- 公開入口 --------
    def process(self):
        wb = load_workbook(self.file_path)
        ws = wb[wb.sheetnames[0]]
        self._hide_rows_with_dash_in_M_to_S(ws, start_row=self.DATA_ROW_START, end_row=self.DATA_ROW_END_FALLBACK)
        
        
        # a) 清乾淨模板/舊檔殘留
        self._wipe_page_setup(ws)

        # b) 取得列印範圍的最後一列（只看 A..U）
        last_row = self._last_filled_row_in_AU(ws) or self.PRINT_LAST_ROW_FALLBACK

        # c) 設定列印範圍
        ws.print_area = f"{self.LEFT_COL}1:{self.RIGHT_COL}{last_row}"

        # d) 每頁重複表頭（可選）
        if self.REPEAT_TITLE_ROWS:
            ws.print_title_rows = self.REPEAT_TITLE_ROWS

        # e) 一頁寬、高度不限（保留手動分頁）
        self._fit_one_page_wide(ws)

        # f) 找「花卉(把)」列 → 在它前一列放手動分頁（讓它成為第二頁第一列）
        target_row = self._find_row_contains(ws, self.TARGET_TEXT,
                                             min_row=self.SEARCH_START_ROW,
                                             max_row=last_row,
                                             max_col=self.SEARCH_MAX_COL)
        if target_row and target_row > 1:
            # Break(id=N) 的語意：在 N 與 N+1 之間換頁
            ws.row_breaks = ws.row_breaks.__class__()  # 先清既有 breaks
            ws.row_breaks.append(Break(id=target_row - 1))

        wb.save(self.file_path)

    # -------- 下面是內部工具 --------
    def _wipe_page_setup(self, ws):
        """把列印相關設定、分頁等清乾淨（不動內容）。"""
        # 列印區與標題列/欄
        try: ws.print_area = None
        except Exception: pass
        try: ws.print_title_rows = None
        except Exception: pass
        try: ws.print_title_cols = None
        except Exception: pass

        # 關閉 fit-to-page 旗標；把 fit 參數與縮放回到中性值
        ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=False)
        ps = ws.page_setup
        ps.scale = 100
        ps.fitToWidth = 0
        ps.fitToHeight = 0
        ps.orientation = "portrait"   # 直向

        # 清掉手動分頁
        ws.row_breaks = ws.row_breaks.__class__()  # reset
        ws.column_breaks = ws.column_breaks.__class__()  # reset

    def _fit_one_page_wide(self, ws):
        """
        開啟以頁數調整模式，寬度 1 頁、高度不限。
        注意：要讓手動分頁生效，fitToHeight 要設 0。
        """
        ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
        ps = ws.page_setup
        ps.scale = None          # 需清掉 scale，fit 才會生效
        ps.fitToWidth = 1        # 一頁寬
        ps.fitToHeight = 0       # 高度不限（尊重手動分頁）
        ps.orientation = "portrait"

    def _last_filled_row_in_AU(self, ws) -> int:
        """回傳 A..U 區間最後一列（若沒資料回 0）。"""
        last = 0
        max_r = ws.max_row or 0
        for r in range(1, max_r + 1):
            for c in range(1, self.SEARCH_MAX_COL + 1):  # A..U
                v = ws.cell(row=r, column=c).value
                if v not in (None, ""):
                    last = max(last, r)
        return last

    def _find_row_contains(self, ws, text: str, min_row: int, max_row: int, max_col: int) -> int | None:
        """在 A..U、min_row..max_row 範圍內找包含指定文字的列，回傳列號。"""
        target = str(text)
        for r in range(min_row, max_row + 1):
            for c in range(1, max_col + 1):
                v = ws.cell(row=r, column=c).value
                if v is None:
                    continue
                # 精準比對可改成：str(v).strip() == target
                if target in str(v).strip():
                    return r
        return None


    def _hide_rows_with_dash_in_M_to_S(self, ws, start_row: int, end_row: int):
        end_row = min(end_row, ws.max_row or end_row)
        for r in range(start_row, end_row + 1):
            vals = []
            for c in range(self.COL_M, self.COL_S + 1):
                v = ws.cell(row=r, column=c).value
                vals.append((str(v).strip() if v is not None else ""))
            if vals and all(v == "-" for v in vals):
                ws.row_dimensions[r].hidden = True